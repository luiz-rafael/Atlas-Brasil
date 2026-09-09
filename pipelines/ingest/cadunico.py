#!/usr/bin/env python3
"""
Ingest CadÚnico — famílias cadastradas / baixa renda por município.

Ordem:
  1) ATLAS_CADUNICO_FILE (csv/xlsx)
  2) SAGI misocial (Solr) — cadun_qtd_familias_* por município/mês
  3) Ipeadata OData (Metadados/Valores) — séries cadunico/bolsa se UF/mun
  4) CKAN dados.gov.br (quando autenticável) + URLs MDS legadas
  5) SKIPPED exit 0

Bronze: cadunico_extract.jsonl
Indicadores: ind_cadunico_familias, ind_cadunico_baixa_renda
"""

from __future__ import annotations

import csv
import io
import os
import re
import subprocess
import sys
from pathlib import Path

import httpx

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from pipelines.common import bronze_dir, day_stamp, throttle, utc_now, write_json, write_jsonl  # noqa: E402
from pipelines.registry import mark_ingested, start_run  # noqa: E402

UA = os.getenv("ATLAS_USER_AGENT", "ATLAS-BRASIL-Ingestor/5.1 (+pesquisa CadUnico)")
SAGI_SOLR = "https://aplicacoes.mds.gov.br/sagi/servicos/misocial"
IPEA_META = "http://www.ipeadata.gov.br/api/odata4/Metadados"
IPEA_VAL = "http://www.ipeadata.gov.br/api/odata4/ValoresSerie"
CKAN_HOSTS = [
    "https://dados.gov.br/api/3/action/package_search",
    "https://dados.gov.br/dados/api/3/action/package_search",
]

CANDIDATE_URLS = [
    "https://aplicacoes.mds.gov.br/sagi/dados/misocial/cadunico_familia_municipio.csv",
]


def _norm(s: str) -> str:
    import unicodedata

    t = unicodedata.normalize("NFKD", s or "")
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", t.lower())


def download_bytes(url: str, dest: Path) -> dict:
    if not any(url.lower().endswith(ext) for ext in (".csv", ".xlsx", ".zip", ".xls")):
        return {"url": url, "error": "não é arquivo tabular"}
    throttle()
    dest.parent.mkdir(parents=True, exist_ok=True)
    try:
        r = httpx.get(url, headers={"User-Agent": UA}, timeout=180.0, follow_redirects=True)
        if r.status_code == 200 and r.content and len(r.content) > 200:
            dest.write_bytes(r.content)
            return {"url": url, "file": dest.name, "bytes": len(r.content), "via": "httpx"}
        raise RuntimeError(f"HTTP {r.status_code}")
    except Exception as first:
        cmd = [
            "curl.exe" if sys.platform == "win32" else "curl",
            "-L",
            "--retry",
            "2",
            "-A",
            UA,
            "-o",
            str(dest),
            url,
        ]
        try:
            subprocess.run(cmd, check=True, timeout=300)
            if dest.exists() and dest.stat().st_size > 200:
                return {"url": url, "file": dest.name, "bytes": dest.stat().st_size, "via": "curl"}
        except Exception as second:
            return {"url": url, "error": f"{first}; curl: {second}"}
        return {"url": url, "error": str(first)}


def _parse_num(v) -> float | None:
    if v is None:
        return None
    s = str(v).strip()
    if not s or s in ("-", "NA", "nan"):
        return None
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _pick(fields: list[str], aliases: list[str]) -> str | None:
    norms = {_norm(f): f for f in fields}
    for a in aliases:
        na = _norm(a)
        if na in norms:
            return norms[na]
        for nk, orig in norms.items():
            if na and len(na) >= 5 and na in nk:
                return orig
    return None


def _ibge_code(raw) -> str | None:
    code = re.sub(r"\D", "", str(raw or ""))
    if len(code) < 6:
        return None
    return code[:7] if len(code) >= 7 else code[:6]


def parse_csv_text(text: str) -> list[dict]:
    delim = ";" if text[:4096].count(";") >= text[:4096].count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    if not reader.fieldnames:
        return []
    fields = list(reader.fieldnames)
    col_ibge = _pick(fields, ["cod_ibge", "codigo_ibge", "ibge", "cd_ibge", "co_municipio"])
    col_ano = _pick(fields, ["ano", "ano_referencia", "nu_ano", "referencia"])
    col_fam = _pick(
        fields,
        ["familias", "qtd_familias", "familias_cadastradas", "total_familias", "qt_familias"],
    )
    col_baixa = _pick(
        fields,
        [
            "baixa_renda",
            "familias_baixa_renda",
            "pobres",
            "extrema_pobreza",
            "familias_pobreza",
            "qt_baixa_renda",
        ],
    )
    if not col_ibge or (not col_fam and not col_baixa):
        return []
    year_default = int(os.getenv("ATLAS_CADUNICO_YEAR", "2023"))
    now = utc_now()
    out = []
    for row in reader:
        code = _ibge_code(row.get(col_ibge))
        if not code:
            continue
        year = year_default
        if col_ano and row.get(col_ano):
            ys = re.sub(r"\D", "", str(row[col_ano]))[:4]
            if ys.isdigit():
                year = int(ys)
        fam = _parse_num(row.get(col_fam)) if col_fam else None
        baixa = _parse_num(row.get(col_baixa)) if col_baixa else None
        if fam is None and baixa is None:
            continue
        out.append(
            {
                "nivel": "MUNICIPALITY",
                "territory_id": f"mun_{code}",
                "cod_ibge": code,
                "ano": year,
                "familias": fam,
                "baixa_renda": baixa,
                "retrieved_at": now,
                "fonte": "cadunico",
            }
        )
    return out


def parse_file(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        return parse_csv_text(path.read_text(encoding="utf-8", errors="replace"))
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            return []
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sh = wb[wb.sheetnames[0]]
        rows = list(sh.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(c or "") for c in rows[0]]
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(header)
        for r in rows[1:]:
            w.writerow(list(r))
        return parse_csv_text(buf.getvalue())
    return []


def fetch_sagi_misocial() -> tuple[list[dict], dict]:
    """Solr MDS/SAGI — famílias CadÚnico por município no mês mais recente com dados."""
    meta: dict = {"via": "sagi_misocial", "url": SAGI_SOLR}
    fl = (
        "anomes_s,codigo_ibge,sigla_uf,municipio,"
        "cadun_qtd_familias_cadastradas_i,cadun_qtd_familias_cadastradas_baixa_renda_i"
    )
    try:
        throttle()
        r = httpx.get(
            SAGI_SOLR,
            params={
                "wt": "json",
                "q": "cadun_qtd_familias_cadastradas_i:[1 TO *]",
                "rows": "1",
                "sort": "anomes_s desc",
                "fl": "anomes_s",
            },
            headers={"User-Agent": UA},
            timeout=90,
        )
        r.raise_for_status()
        docs = (r.json().get("response") or {}).get("docs") or []
        if not docs:
            meta["error"] = "nenhum doc com cadun_qtd_familias_cadastradas_i"
            return [], meta
        anomes = str(docs[0].get("anomes_s") or "")
        if len(anomes) < 6:
            meta["error"] = f"anomes inválido: {anomes}"
            return [], meta
        year = int(anomes[:4])
        meta["anomes"] = anomes
        meta["ano"] = year

        rows: list[dict] = []
        now = utc_now()
        start = 0
        page = 2000
        total = None
        while True:
            throttle()
            rr = httpx.get(
                SAGI_SOLR,
                params={
                    "wt": "json",
                    "q": f"anomes_s:{anomes} AND cadun_qtd_familias_cadastradas_i:[* TO *]",
                    "rows": str(page),
                    "start": str(start),
                    "fl": fl,
                },
                headers={"User-Agent": UA},
                timeout=120,
            )
            rr.raise_for_status()
            resp = rr.json().get("response") or {}
            if total is None:
                total = int(resp.get("numFound") or 0)
                meta["numFound"] = total
            batch = resp.get("docs") or []
            if not batch:
                break
            for d in batch:
                code = _ibge_code(d.get("codigo_ibge"))
                if not code:
                    continue
                fam = _parse_num(d.get("cadun_qtd_familias_cadastradas_i"))
                baixa = _parse_num(d.get("cadun_qtd_familias_cadastradas_baixa_renda_i"))
                if fam is None and baixa is None:
                    continue
                uf = str(d.get("sigla_uf") or "").upper()[:2]
                rows.append(
                    {
                        "nivel": "MUNICIPALITY",
                        "territory_id": f"mun_{code}",
                        "cod_ibge": code,
                        "uf": uf,
                        "ano": year,
                        "anomes": anomes,
                        "familias": fam,
                        "baixa_renda": baixa,
                        "retrieved_at": now,
                        "fonte": "sagi_misocial",
                    }
                )
            start += len(batch)
            print(f"  SAGI misocial {anomes}: {start}/{total}", flush=True)
            if start >= total or len(batch) < page:
                break
        meta["rows"] = len(rows)
        return rows, meta
    except Exception as e:
        meta["error"] = str(e)
        return [], meta


def fetch_ipeadata() -> tuple[list[dict], dict]:
    """Fallback: séries Ipeadata com cadunico/bolsa em nível Municípios/Estados."""
    meta: dict = {"via": "ipeadata", "url": IPEA_META}
    try:
        throttle()
        r = httpx.get(IPEA_META, headers={"User-Agent": UA}, timeout=180)
        r.raise_for_status()
        vals = r.json().get("value") or []
        keys = ("cadun", "cadastro unico", "cadastro único", "bolsa fam")
        # Preferir séries cujo nome cite cadastro/famílias; PBF é proxy fraco — só se nome explícito cadunico
        candidates = []
        for it in vals:
            blob = " ".join(
                str(it.get(k) or "") for k in ("SERCODIGO", "SERNOME", "SERCOMENTARIO")
            ).lower()
            if "cadun" in blob or "cadastro unico" in blob or "cadastro único" in blob:
                candidates.append(it)
        meta["candidates"] = [
            {"SERCODIGO": c.get("SERCODIGO"), "SERNOME": c.get("SERNOME"), "NIVNOME": c.get("NIVNOME")}
            for c in candidates[:10]
        ]
        if not candidates:
            meta["error"] = "sem série CadÚnico no Metadados Ipeadata"
            return [], meta

        now = utc_now()
        out: list[dict] = []
        for it in candidates[:3]:
            code = it.get("SERCODIGO")
            if not code:
                continue
            throttle()
            vr = httpx.get(
                f"{IPEA_VAL}(SERCODIGO='{code}')",
                headers={"User-Agent": UA},
                timeout=300,
            )
            if vr.status_code != 200:
                continue
            items = vr.json().get("value") or []
            # último ano por território
            latest: dict[str, dict] = {}
            for row in items:
                niv = str(row.get("NIVNOME") or "")
                ter = str(row.get("TERCODIGO") or "")
                val = row.get("VALVALOR")
                if val is None or not ter:
                    continue
                ys = str(row.get("VALDATA") or "")[:4]
                if not ys.isdigit():
                    continue
                year = int(ys)
                if "munic" in niv.lower():
                    ibge = _ibge_code(ter)
                    if not ibge:
                        continue
                    key = f"mun_{ibge}"
                    nivel, tid, cod = "MUNICIPALITY", key, ibge
                elif "estado" in niv.lower() or niv.lower() == "estados":
                    # TERCODIGO UF numérico IBGE 2 dígitos — guardar como código
                    uf_code = re.sub(r"\D", "", ter)
                    if len(uf_code) != 2:
                        continue
                    key = f"ufcode_{uf_code}"
                    nivel, tid, cod = "STATE", f"uf_{uf_code}", uf_code
                else:
                    continue
                prev = latest.get(key)
                if prev is None or year >= prev["ano"]:
                    latest[key] = {
                        "nivel": nivel,
                        "territory_id": tid,
                        "cod_ibge": cod,
                        "ano": year,
                        "familias": float(val),
                        "baixa_renda": None,
                        "retrieved_at": now,
                        "fonte": f"ipeadata:{code}",
                    }
            out.extend(latest.values())
            if out:
                meta["serie"] = code
                break
        meta["rows"] = len(out)
        if not out:
            meta["error"] = "séries CadÚnico sem valores territoriais"
        return out, meta
    except Exception as e:
        meta["error"] = str(e)
        return [], meta


def try_ckan(outdir: Path) -> list[dict]:
    metas = []
    for host in CKAN_HOSTS:
        try:
            throttle()
            r = httpx.get(
                host,
                params={"q": "cadunico OR \"cadastro unico\" familias", "rows": 8},
                headers={"User-Agent": UA},
                timeout=45,
            )
            if r.status_code != 200:
                metas.append({"ckan": host, "error": f"HTTP {r.status_code}"})
                continue
            for pkg in (r.json().get("result") or {}).get("results") or []:
                for res in pkg.get("resources") or []:
                    url = res.get("url") or ""
                    fmt = (res.get("format") or "").lower()
                    if not url:
                        continue
                    if fmt in ("csv", "xlsx", "xls", "zip") or any(
                        url.lower().endswith(ext) for ext in (".csv", ".xlsx", ".xls", ".zip")
                    ):
                        dest = outdir / f"ckan_{Path(url.split('?')[0]).name}"
                        metas.append(download_bytes(url, dest))
        except Exception as e:
            metas.append({"ckan": host, "error": str(e)})
    return metas


def main() -> int:
    run = start_run("cadunico", "cadunico.familias")
    out = bronze_dir("cadunico")
    raw = out / "raw"
    raw.mkdir(exist_ok=True)
    rows: list[dict] = []
    files_meta: list[dict] = []
    messages: list[str] = []

    local = os.getenv("ATLAS_CADUNICO_FILE")
    if local:
        p = Path(local)
        if p.is_file():
            rows = parse_file(p)
            files_meta.append({"file": str(p), "via": "ATLAS_CADUNICO_FILE", "rows": len(rows)})
        else:
            messages.append(f"ATLAS_CADUNICO_FILE ausente: {local}")

    if not rows:
        got, meta = fetch_sagi_misocial()
        files_meta.append(meta)
        if got:
            rows = got
            messages.append(f"SAGI misocial anomes={meta.get('anomes')} rows={len(got)}")

    if not rows:
        got, meta = fetch_ipeadata()
        files_meta.append(meta)
        if got:
            rows = got
            messages.append(f"Ipeadata serie={meta.get('serie')} rows={len(got)}")

    if not rows:
        for url in CANDIDATE_URLS:
            dest = raw / Path(url.split("?")[0]).name
            if not dest.suffix:
                continue
            meta = download_bytes(url, dest)
            files_meta.append(meta)
            if meta.get("error"):
                continue
            got = parse_file(dest)
            if got:
                rows = got
                break

    if not rows:
        files_meta.extend(try_ckan(raw))
        for m in files_meta:
            if m.get("file") and not m.get("error"):
                p = raw / m["file"]
                if p.exists():
                    got = parse_file(p)
                    if got:
                        rows = got
                        break

    if rows:
        write_jsonl(out / "cadunico_extract.jsonl", rows)
    else:
        messages.append(
            "SKIPPED cadunico: sem CSV/API pública utilizável. "
            "Defina ATLAS_CADUNICO_FILE=caminho.csv (IBGE + familias [, baixa_renda])."
        )
        print(messages[-1], flush=True)

    write_json(
        out / "cadunico_meta.json",
        {
            "em": utc_now(),
            "fonte": "cadunico",
            "rows": len(rows),
            "files": files_meta,
            "messages": messages,
            "ingestion_run_id": run["ingestion_run_id"],
            "day": day_stamp(),
            "portal": "https://aplicacoes.mds.gov.br/sagi/servicos/misocial",
        },
    )
    mark_ingested(
        "cadunico",
        run_id=run["ingestion_run_id"],
        counts={"rows": len(rows)},
        ok=len(rows) > 0,
        error=None if rows else (messages[-1] if messages else "sem dados"),
        dataset_id="cadunico.familias",
    )
    print(f"OK CadÚnico: rows={len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
