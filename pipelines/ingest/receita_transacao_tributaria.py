#!/usr/bin/env python3
"""
Receita Federal — Transações tributárias (acordos administrativos).

Editais públicos frequentemente não trazem CSV row-level.
Se só HTML/PDF → SKIPPED exit 0 + ATLAS_TRANSACAO_FILE.
Eventos = ADMINISTRATIVE settlement — NÃO LEGAL_CASE.
"""

from __future__ import annotations

import csv
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from pipelines.common import (  # noqa: E402
    append_event,
    bronze_dir,
    utc_now,
    write_json,
    write_jsonl,
    write_manifest,
    write_raw_record,
)
from pipelines.ingest.receita_common import (  # noqa: E402
    company_id_from_cnpj,
    copy_env_file,
    load_dotenv,
    only_cnpj14,
    scrape_download_links,
    to_float,
    try_urls,
)
from pipelines.registry import mark_ingested, start_run  # noqa: E402

SOURCE_ID = "receita_transacao_tributaria"
DATASET_ID = "rfb.transacao_tributaria"
PORTAL = "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/editais/transacao-tributaria"
CANDIDATE_PAGES = [
    PORTAL,
    "https://www.gov.br/receitafederal/pt-br/assuntos/orientacao-tributaria/transacao-tributaria",
    "https://www.gov.br/receitafederal/pt-br/acesso-a-informacao/dados-abertos",
]


def probe_structured_urls() -> list[str]:
    urls: list[str] = []
    for page in CANDIDATE_PAGES:
        try:
            found = scrape_download_links(
                page,
                keywords=("csv", "xlsx", "ods", "zip", "transacao", "adesao"),
            )
            for u in found:
                low = u.lower()
                # ignore pure PDFs of editais
                if low.endswith(".pdf") or low.endswith(".pdf/view"):
                    continue
                if any(ext in low for ext in (".csv", ".xlsx", ".ods", ".zip", "@@download")):
                    if u.endswith("/view"):
                        u = u[: -len("/view")] + "/@@download/file"
                    urls.append(u)
        except Exception:
            continue
    seen: set[str] = set()
    out: list[str] = []
    for u in urls:
        if u not in seen:
            seen.add(u)
            out.append(u)
    return out


def parse_structured(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    rows: list[dict] = []
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8", errors="ignore")
        sample = text[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        for i, row in enumerate(reader):
            rows.append(_normalize_row(row, i, path.name))
        return rows
    if suffix in (".xlsx", ".xls", ".ods"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        headers: list[str] | None = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            cells = ["" if c is None else str(c).strip() for c in row]
            if headers is None:
                if any(cells):
                    headers = [c.lower() for c in cells]
                continue
            d = {headers[j]: cells[j] if j < len(cells) else "" for j in range(len(headers))}
            rows.append(_normalize_row(d, i, path.name))
        wb.close()
        return rows
    raise ValueError(f"formato não estruturado: {suffix}")


def _normalize_row(row: dict, idx: int, source_file: str) -> dict:
    # flexible key search
    flat = {str(k).lower(): v for k, v in row.items()}
    cnpj = None
    for k, v in flat.items():
        if "cnpj" in k:
            cnpj = only_cnpj14(v)
            if cnpj:
                break
    valor = None
    for k, v in flat.items():
        if any(x in k for x in ("valor", "montante", "credito", "débito", "debito")):
            valor = to_float(v)
            if valor is not None:
                break
    edital = None
    for k, v in flat.items():
        if "edital" in k or "numero" in k:
            edital = str(v).strip() or None
            if edital:
                break
    return {
        "event_id": f"tax_tx_{idx}_{cnpj or 'na'}",
        "event_type": "ADMINISTRATIVE_TAX_SETTLEMENT",
        "edital": edital,
        "cnpj": cnpj,
        "company_id": company_id_from_cnpj(cnpj),
        "valor": valor,
        "raw": {k: (str(v)[:200] if v is not None else None) for k, v in list(row.items())[:30]},
        "fonte": SOURCE_ID,
        "dataset_id": DATASET_ID,
        "nota": "Acordo administrativo — não é LEGAL_CASE / DataJud",
        "source_file": source_file,
    }


def try_download(out: Path) -> tuple[Path | None, dict]:
    env_path, env_meta = copy_env_file("ATLAS_TRANSACAO_FILE", out)
    if env_path:
        return env_path, env_meta

    urls = probe_structured_urls()
    if not urls:
        return None, {
            "error": "apenas editais HTML/PDF — sem CSV/XLSX estruturado detectado",
            "portal": PORTAL,
            "candidates": [],
        }
    path, meta = try_urls(urls[:10], out, filename="transacao_tributaria_data.bin", min_bytes=200)
    return path, {**meta, "probed_urls": urls[:10]}


def main() -> int:
    load_dotenv()
    try:
        run = start_run(SOURCE_ID, DATASET_ID)
    except Exception as e:
        print(f"fail-soft start_run {SOURCE_ID}: {e}", file=sys.stderr)
        return 0

    out = bronze_dir(SOURCE_ID)
    file_path, meta_dl = try_download(out)
    meta = {
        "fetched_at": utc_now(),
        "portal": PORTAL,
        "ingestion_run_id": run["ingestion_run_id"],
        "entity": "ADMINISTRATIVE_TAX_SETTLEMENT",
        "nota": "Não mesclar com LEGAL_CASE / DataJud",
        **meta_dl,
    }

    if not file_path:
        meta["status"] = "SKIPPED"
        meta["reason"] = (
            "Sem dataset estruturado público (editais HTML/PDF). "
            "Defina ATLAS_TRANSACAO_FILE com CSV/XLSX local se disponível."
        )
        write_json(out / "meta.json", meta)
        write_manifest(out, SOURCE_ID, [], extra=meta)
        mark_ingested(SOURCE_ID, run_id=run["ingestion_run_id"], counts={"rows": 0}, dataset_id=DATASET_ID, ok=True)
        print(
            "SKIPPED receita_transacao_tributaria: só editais HTML/PDF — "
            "use ATLAS_TRANSACAO_FILE para CSV/XLSX estruturado.",
            flush=True,
        )
        return 0

    try:
        # rename by sniff
        raw = file_path.read_bytes()
        if raw[:2] == b"PK":
            named = out / "transacao.xlsx"
        elif b"," in raw[:200] or b";" in raw[:200]:
            named = out / "transacao.csv"
        else:
            named = file_path
        if named != file_path:
            named.write_bytes(raw)
            file_path = named
        rows = parse_structured(file_path)
    except Exception as e:
        meta["status"] = "SKIPPED"
        meta["parse_error"] = str(e)
        write_json(out / "meta.json", meta)
        mark_ingested(SOURCE_ID, run_id=run["ingestion_run_id"], counts={"rows": 0}, dataset_id=DATASET_ID, ok=True)
        print(f"SKIPPED {SOURCE_ID} parse: {e}", flush=True)
        return 0

    write_jsonl(out / "tax_transaction_extract.jsonl", rows)
    write_json(out / "meta.json", {**meta, "rows": len(rows), "status": "OK"})
    write_raw_record(
        source_id=SOURCE_ID,
        ingestion_run_id=run["ingestion_run_id"],
        connector_version=run["connector_version"],
        payload=file_path.read_bytes(),
        filename=file_path.name,
        source_url=PORTAL,
        dataset_id=DATASET_ID,
    )
    write_manifest(out, SOURCE_ID, [{"file": file_path.name, "bytes": file_path.stat().st_size}], extra=meta)
    append_event("document.discovered", {"source": SOURCE_ID, "rows": len(rows)})
    mark_ingested(
        SOURCE_ID,
        run_id=run["ingestion_run_id"],
        counts={"rows": len(rows)},
        dataset_id=DATASET_ID,
        ok=True,
    )
    print(f"OK {SOURCE_ID} rows={len(rows)} -> {out}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
