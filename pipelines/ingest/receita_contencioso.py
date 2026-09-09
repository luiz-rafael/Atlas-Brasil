#!/usr/bin/env python3
"""
Receita Federal — Contencioso administrativo tributário.

Entidade: ADMINISTRATIVE_TAX_CASE / indicadores de acervo.
NUNCA mesclar com DataJud LEGAL_CASE.
"""

from __future__ import annotations

import csv
import io
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from pipelines.common import (  # noqa: E402
    append_event,
    bronze_dir,
    http_get,
    utc_now,
    write_json,
    write_jsonl,
    write_manifest,
    write_raw_record,
)
from pipelines.ingest.receita_common import (  # noqa: E402
    load_dotenv,
    scrape_download_links,
    to_float,
    try_urls,
    year_from_header,
)
from pipelines.registry import mark_ingested, start_run  # noqa: E402

SOURCE_ID = "receita_contencioso"
DATASET_ID = "rfb.contencioso_admin"
PORTAL = (
    "https://www.gov.br/receitafederal/pt-br/acesso-a-informacao/"
    "dados-abertos/contencioso-administrativo"
)
CKAN_URLS = [
    "https://dados.gov.br/api/3/action/package_show?id=contencioso-administrativo-de-primeira-instancia-e-de-segunda-instancia-na-rfb",
    "https://dados.gov.br/api/3/action/package_search?q=contencioso+administrativo+receita",
]
# Repositório RFB /dados (CKAN dados.gov frequentemente 401)
KNOWN_CSV_URLS = [
    "https://www.gov.br/receitafederal/dados/contencioso-administrativo-de-primeira-instancia.csv/@@download/file",
    "https://www.gov.br/receitafederal/dados/contencioso-administrativo-de-primeira-instancia-pequeno-valor.csv/@@download/file",
    "https://www.gov.br/receitafederal/dados/contencioso-administrativo-de-segunda-instancia.csv/@@download/file",
]


def _year_from_mes_ano(val: object) -> int | None:
    """Aceita 'jan-13', 'dez/2024', '2023-01' e anos completos."""
    if val is None:
        return None
    s = str(val).strip().lower()
    m = re.search(r"(?:jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)[a-z]*[-/](\d{2,4})", s)
    if m:
        yy = int(m.group(1))
        if yy < 100:
            yy += 2000 if yy < 80 else 1900
        return yy if 1990 <= yy <= 2100 else None
    y = year_from_header(val)
    return y


def probe_ckan() -> list[str]:
    urls: list[str] = []
    for api in CKAN_URLS:
        try:
            r = http_get(api, timeout=60.0)
            if r.status_code != 200 or not r.content:
                continue
            blob = r.json()
            result = blob.get("result") or {}
            resources = result.get("resources") or []
            if not resources and "results" in result:
                for pkg in result.get("results") or []:
                    resources.extend(pkg.get("resources") or [])
            for res in resources:
                u = res.get("url") or res.get("download_url")
                if u:
                    urls.append(u)
        except Exception:
            continue
    return urls


def probe_portal() -> list[str]:
    urls: list[str] = []
    pages = [
        PORTAL,
        f"{PORTAL}/contencioso-administrativo",
        "https://www.gov.br/receitafederal/dados",
    ]
    for page in pages:
        try:
            found = scrape_download_links(
                page,
                keywords=("contencioso", "acervo"),
            )
            for u in found:
                low = u.lower()
                if "contencioso" not in low and "acervo" not in low:
                    continue
                if low.endswith(".pdf") or low.endswith(".pdf/view"):
                    continue
                if u.endswith("/view"):
                    u = u[: -len("/view")] + "/@@download/file"
                urls.append(u)
        except Exception:
            continue
    # também varre /dados só por nomes com "contencioso"
    try:
        r = http_get("https://www.gov.br/receitafederal/dados", timeout=90.0)
        if r.status_code == 200:
            for m in re.findall(r'href=["\']([^"\']*contencioso[^"\']*)["\']', r.text, flags=re.I):
                u = m
                if not u.startswith("http"):
                    u = "https://www.gov.br" + u if u.startswith("/") else f"https://www.gov.br/receitafederal/{u}"
                if u.endswith("/view"):
                    u = u[: -len("/view")] + "/@@download/file"
                urls.append(u)
    except Exception:
        pass
    seen: set[str] = set()
    out = []
    for u in urls:
        if u not in seen:
            seen.add(u)
            out.append(u)
    return out


def parse_file(path: Path) -> tuple[list[dict], list[dict]]:
    """Returns (case_or_detail_rows, indicator_rows)."""
    suffix = path.suffix.lower()
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        rows = data if isinstance(data, list) else data.get("rows") or data.get("data") or []
        return [_norm_case(r, i, path.name) for i, r in enumerate(rows)], []
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8", errors="ignore")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        cases, inds = [], []
        for i, row in enumerate(reader):
            c, ind = _classify_row(row, i, path.name)
            if c:
                cases.append(c)
            if ind:
                inds.append(ind)
        return cases, inds
    if suffix in (".xlsx", ".xls", ".ods"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        cases: list[dict] = []
        inds: list[dict] = []
        for sn in wb.sheetnames:
            ws = wb[sn]
            headers = None
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = ["" if c is None else str(c).strip() for c in row]
                if headers is None:
                    if any(cells):
                        headers = [c.lower() for c in cells]
                    continue
                d = {headers[j]: cells[j] if j < len(cells) else "" for j in range(len(headers))}
                c, ind = _classify_row(d, i, path.name)
                if c:
                    cases.append(c)
                if ind:
                    inds.append(ind)
        wb.close()
        return cases, inds
    raise ValueError(f"formato não suportado: {suffix}")


def _classify_row(row: dict, idx: int, source_file: str) -> tuple[dict | None, dict | None]:
    flat = {str(k).lower(): v for k, v in row.items()}
    year = None
    for k, v in flat.items():
        if any(x in k for x in ("mês-ano", "mes-ano", "ano", "year", "compet", "periodo", "período")):
            y = _year_from_mes_ano(v)
            if y:
                year = y
                break
    if year is None:
        for v in flat.values():
            y = _year_from_mes_ano(v)
            if y:
                year = y
                break

    qtd = None
    valor = None
    tempo = None
    for k, v in flat.items():
        kl = k.lower()
        if qtd is None and any(x in kl for x in ("quantidade de processos", "qtd", "quant", "estoque", "acervo")):
            if "julgad" not in kl:
                qtd = to_float(v)
        if valor is None and kl.startswith("valor total dos processos") and "julgad" not in kl:
            valor = to_float(v)
        if tempo is None and "tempo médio" in kl and "todos os processos" in kl and "priorit" not in kl:
            tempo = to_float(v)

    # ID de processo individual — NÃO confundir com "quantidade de processos"
    case_id = None
    for k, v in flat.items():
        kl = k.lower().strip()
        if any(x in kl for x in ("quantidade", "valor", "tempo", "hora", "publica")):
            continue
        if kl in ("npu", "numero", "número", "processo", "id") or "npu" in kl or "numero do processo" in kl:
            case_id = str(v).strip() or None
            if case_id:
                break

    ind = None
    if year and (qtd is not None or valor is not None or tempo is not None) and not case_id:
        periodo = None
        for k, v in flat.items():
            if "mês-ano" in k or "mes-ano" in k:
                periodo = str(v).strip()
                break
        ind = {
            "year": year,
            "period_label": periodo,
            "quantidade": qtd,
            "valor": valor,
            "tempo_medio": tempo,
            "entity_type": "ADMINISTRATIVE_TAX_CASE_STOCK",
            "fonte": SOURCE_ID,
            "dataset_id": DATASET_ID,
            "source_file": source_file,
            "nota": "Acervo administrativo — não é LEGAL_CASE",
        }

    case = None
    if case_id:
        case = _norm_case(row, idx, source_file)
        case["admin_case_id"] = case_id
    return case, ind


def _norm_case(row: dict, idx: int, source_file: str) -> dict:
    return {
        "admin_case_id": f"admin_tax_{idx}",
        "entity_type": "ADMINISTRATIVE_TAX_CASE",
        "raw": {str(k): (str(v)[:300] if v is not None else None) for k, v in list(row.items())[:40]},
        "fonte": SOURCE_ID,
        "dataset_id": DATASET_ID,
        "source_file": source_file,
        "nota": "ADMINISTRATIVE_TAX_CASE — separado de DataJud LEGAL_CASE",
    }


def _resolve_local_inputs(out: Path) -> tuple[list[Path], dict]:
    """ATLAS_CONTENCIOSO_FILE = arquivo, vários (;), ou diretório com CSVs."""
    import os
    from shutil import copy2

    raw = os.getenv("ATLAS_CONTENCIOSO_FILE", "").strip()
    if not raw:
        return [], {}
    found: list[Path] = []
    for part in raw.split(";"):
        p = Path(part.strip())
        if not p.is_absolute():
            p = (ROOT / p).resolve()
        if p.is_dir():
            found.extend(sorted(p.glob("contencioso*.csv")))
            if not found:
                found.extend(sorted(p.glob("*.csv")))
        elif p.is_file():
            found.append(p)
    if not found:
        return [], {"via": "ATLAS_CONTENCIOSO_FILE", "error": f"nada em {raw}"}
    copied: list[Path] = []
    for src in found:
        dest = out / src.name
        copy2(src, dest)
        copied.append(dest)
    return copied, {
        "via": "ATLAS_CONTENCIOSO_FILE",
        "files": [p.name for p in copied],
        "bytes": sum(p.stat().st_size for p in copied),
    }


def try_download(out: Path) -> tuple[list[Path], dict]:
    local, env_meta = _resolve_local_inputs(out)
    if local:
        return local, env_meta

    # Preferência: CSVs oficiais no repositório RFB
    from pipelines.ingest.receita_common import download_to

    saved: list[Path] = []
    attempts: list[dict] = []
    for i, url in enumerate(KNOWN_CSV_URLS):
        name = {
            0: "contencioso_1a.csv",
            1: "contencioso_1a_pequeno_valor.csv",
            2: "contencioso_2a.csv",
        }.get(i, f"contencioso_{i}.csv")
        path, meta = download_to(url, out / name, timeout=120.0, min_bytes=200)
        attempts.append(meta)
        if path:
            saved.append(path)
    if saved:
        return saved, {"via": "rfb_dados", "files": [p.name for p in saved], "attempts": attempts}

    urls = probe_ckan() + probe_portal()
    if urls:
        path, meta = try_urls(
            urls[:12],
            out,
            filename="contencioso_admin.bin",
            min_bytes=200,
            timeout=120.0,
        )
        if path:
            return [path], {**meta, "probed_urls": urls[:12]}
    return [], {
        "error": "CKAN 401/vazio e portal sem CSV/XLSX de contencioso",
        "portal": PORTAL,
        "ckan_tried": CKAN_URLS,
        "known_attempts": attempts,
    }


def main() -> int:
    load_dotenv()
    try:
        run = start_run(SOURCE_ID, DATASET_ID)
    except Exception as e:
        print(f"fail-soft start_run {SOURCE_ID}: {e}", file=sys.stderr)
        return 0

    out = bronze_dir(SOURCE_ID)
    files, meta_dl = try_download(out)
    meta = {
        "fetched_at": utc_now(),
        "portal": PORTAL,
        "ingestion_run_id": run["ingestion_run_id"],
        "entity": "ADMINISTRATIVE_TAX_CASE",
        "nota": "Nunca mesclar com LEGAL_CASE / DataJud",
        **meta_dl,
    }

    if not files:
        meta["status"] = "SKIPPED"
        meta["reason"] = (
            "Sem download utilizável (CKAN/dados.gov frequentemente 401; portal sem arquivo). "
            "Defina ATLAS_CONTENCIOSO_FILE com CSV/XLSX local."
        )
        write_json(out / "meta.json", meta)
        write_manifest(out, SOURCE_ID, [], extra=meta)
        mark_ingested(SOURCE_ID, run_id=run["ingestion_run_id"], counts={"rows": 0}, dataset_id=DATASET_ID, ok=True)
        print(
            "SKIPPED receita_contencioso: sem arquivo. Use ATLAS_CONTENCIOSO_FILE.",
            flush=True,
        )
        return 0

    cases: list[dict] = []
    inds: list[dict] = []
    try:
        for file_path in files:
            c, i = parse_file(file_path)
            cases.extend(c)
            inds.extend(i)
    except Exception as e:
        meta["status"] = "SKIPPED"
        meta["parse_error"] = str(e)
        write_json(out / "meta.json", meta)
        mark_ingested(SOURCE_ID, run_id=run["ingestion_run_id"], counts={"rows": 0}, dataset_id=DATASET_ID, ok=True)
        print(f"SKIPPED {SOURCE_ID} parse: {e}", flush=True)
        return 0

    write_jsonl(out / "admin_tax_contencioso_extract.jsonl", cases)
    write_jsonl(out / "admin_tax_contencioso_indicators.jsonl", inds)
    write_json(
        out / "meta.json",
        {**meta, "cases": len(cases), "indicators": len(inds), "status": "OK"},
    )
    for file_path in files:
        write_raw_record(
            source_id=SOURCE_ID,
            ingestion_run_id=run["ingestion_run_id"],
            connector_version=run["connector_version"],
            payload=file_path.read_bytes(),
            filename=file_path.name,
            source_url=PORTAL,
            dataset_id=DATASET_ID,
        )
    write_manifest(
        out,
        SOURCE_ID,
        [{"file": p.name, "bytes": p.stat().st_size} for p in files],
        extra=meta,
    )
    append_event(
        "document.discovered",
        {"source": SOURCE_ID, "cases": len(cases), "indicators": len(inds)},
    )
    mark_ingested(
        SOURCE_ID,
        run_id=run["ingestion_run_id"],
        counts={"cases": len(cases), "indicators": len(inds)},
        dataset_id=DATASET_ID,
        ok=True,
    )
    print(f"OK {SOURCE_ID} cases={len(cases)} indicators={len(inds)} -> {out}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
