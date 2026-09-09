#!/usr/bin/env python3
"""
Receita Federal — Carga tributária (% PIB).

Indicador separado de ind_arrecadacao_federal.
Metodologia: estudo CTB RFB; nota sobre exclusão FGTS/Sistema S no estudo 2024.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from pipelines.common import (  # noqa: E402
    append_event,
    bronze_dir,
    utc_now,
    write_json,
    write_jsonl,
    write_manifest,
    write_raw_record,
)
from pipelines.ingest.receita_common import (  # noqa: E402
    copy_env_file,
    download_to,
    load_dotenv,
    norm_text,
    to_float,
    year_from_header,
)
from pipelines.registry import mark_ingested, start_run  # noqa: E402

SOURCE_ID = "receita_carga_tributaria"
DATASET_ID = "rfb.carga_tributaria"
PORTAL = (
    "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/"
    "publicacoes/estudos/carga-tributaria"
)
METHODOLOGY_2024 = (
    "Carga Tributária no Brasil (CTB) — RFB. "
    "tax_to_gdp_ratio = arrecadação tributária bruta / PIB. "
    "Estudo 2024: atenção a mudanças metodológicas de cobertura "
    "(FGTS / Sistema S — verificar notas do estudo oficial quanto a inclusão/exclusão "
    "em relação a edições anteriores). "
    "Carga tributária (% PIB) ≠ série de arrecadação federal nominal."
)

YEARS = list(range(2024, 2016, -1))


def candidate_urls() -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    for y in YEARS:
        out.append(
            (
                y,
                (
                    "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/"
                    "publicacoes/estudos/carga-tributaria/"
                    f"tabelas-carga-tributaria-no-brasil-{y}/@@download/file"
                ),
            )
        )
    return out


def parse_tab01(path: Path, study_year: int | None) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows_out: list[dict] = []
    sheet = "Tab_01" if "Tab_01" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]

    header_years: dict[int, int] = {}
    pib_by_year: dict[int, float] = {}
    arrec_by_year: dict[int, float] = {}
    carga_by_year: dict[int, float] = {}
    unit_note = None

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        cells = list(row)
        labels = [norm_text(str(c)) if c is not None else "" for c in cells]
        # unit
        joined = " ".join(labels)
        if "R$ BILHO" in joined or "R$ BILH" in joined:
            unit_note = "R$ bilhões"

        # header with years
        years_found = {
            j: year_from_header(c)
            for j, c in enumerate(cells)
            if year_from_header(c)
        }
        if len(years_found) >= 3 and not header_years:
            # skip label columns
            header_years = {j: y for j, y in years_found.items() if y}
            continue

        if not header_years:
            continue

        label = next((x for x in labels if x), "")
        if not label:
            continue

        bucket = None
        if "PRODUTO INTERNO BRUTO" in label or label == "PIB":
            bucket = pib_by_year
        elif "ARRECADACAO TRIBUTARIA BRUTA" in label:
            bucket = arrec_by_year
        elif "CARGA TRIBUTARIA BRUTA" in label:
            bucket = carga_by_year
        else:
            continue

        for col_i, year in header_years.items():
            if col_i >= len(cells):
                continue
            num = to_float(cells[col_i])
            if num is None:
                continue
            bucket[year] = num

    wb.close()

    years = sorted(set(pib_by_year) | set(arrec_by_year) | set(carga_by_year))
    for year in years:
        ratio = carga_by_year.get(year)
        # ratios sometimes stored as 0.30; sometimes as 30.x
        if ratio is not None and ratio > 1.5:
            ratio = ratio / 100.0
        amount = None
        if year in arrec_by_year:
            # Tab_01 unit is R$ bilhões
            amount = arrec_by_year[year] * 1_000_000_000.0
        rows_out.append(
            {
                "year": year,
                "tax_to_gdp_ratio": ratio,
                "amount": amount,
                "amount_unit": "BRL",
                "pib_bilhoes": pib_by_year.get(year),
                "arrecadacao_bruta_bilhoes": arrec_by_year.get(year),
                "methodology": METHODOLOGY_2024,
                "methodology_notes": {
                    "fgts_sistema_s": (
                        "Estudo 2024 CTB — verificar notas oficiais sobre exclusão/inclusão "
                        "de FGTS e contribuições ao Sistema S vs edições anteriores."
                    ),
                    "unit_sheet": unit_note or "R$ bilhões",
                    "study_file_year": study_year,
                },
                "source": SOURCE_ID,
                "dataset_id": DATASET_ID,
                "source_file": path.name,
                "territory_id": "terr_br",
                "territory_type": "BRASIL",
                "nota": "ind_carga_tributaria_pib ≠ ind_arrecadacao_federal",
            }
        )
    return rows_out


def try_download(out: Path) -> tuple[list[tuple[Path, int | None]], dict]:
    env_path, env_meta = copy_env_file("ATLAS_CARGA_TRIBUTARIA_FILE", out)
    if env_path:
        return [(env_path, None)], env_meta

    files: list[tuple[Path, int | None]] = []
    attempts: list[dict] = []
    for study_year, url in candidate_urls():
        dest = out / f"tabelas_carga_tributaria_{study_year}.xlsx"
        path, meta = download_to(url, dest, timeout=120.0, min_bytes=5000)
        attempts.append(meta)
        if path:
            files.append((path, study_year))
            # prefer latest successful study; still try a couple older for redundancy
            if len(files) >= 2:
                break
    if files:
        return files, {"via": "http", "downloads": attempts, "files": [p.name for p, _ in files]}
    return [], {"error": "download failed", "attempts": attempts}


def main() -> int:
    load_dotenv()
    try:
        run = start_run(SOURCE_ID, DATASET_ID)
    except Exception as e:
        print(f"fail-soft start_run {SOURCE_ID}: {e}", file=sys.stderr)
        return 0

    out = bronze_dir(SOURCE_ID)
    files, meta_dl = try_download(out)
    meta = {
        "fetched_at": utc_now(),
        "portal": PORTAL,
        "ingestion_run_id": run["ingestion_run_id"],
        "methodology": METHODOLOGY_2024,
        **meta_dl,
    }

    if not files:
        meta["status"] = "SKIPPED"
        meta["reason"] = (
            "URL CTB indisponível. Defina ATLAS_CARGA_TRIBUTARIA_FILE com xlsx local."
        )
        write_json(out / "meta.json", meta)
        write_manifest(out, SOURCE_ID, [], extra=meta)
        mark_ingested(SOURCE_ID, run_id=run["ingestion_run_id"], counts={"rows": 0}, dataset_id=DATASET_ID, ok=True)
        print(f"SKIPPED {SOURCE_ID}: sem arquivo. Use ATLAS_CARGA_TRIBUTARIA_FILE.", flush=True)
        return 0

    # Prefer newest study file for the series (it embeds historical years)
    files_sorted = sorted(files, key=lambda x: x[1] or 0, reverse=True)
    primary, study_year = files_sorted[0]
    try:
        rows = parse_tab01(primary, study_year)
    except Exception as e:
        meta["status"] = "SKIPPED"
        meta["parse_error"] = str(e)
        write_json(out / "meta.json", meta)
        mark_ingested(SOURCE_ID, run_id=run["ingestion_run_id"], counts={"rows": 0}, dataset_id=DATASET_ID, ok=True)
        print(f"SKIPPED {SOURCE_ID} parse: {e}", flush=True)
        return 0

    write_jsonl(out / "carga_tributaria_extract.jsonl", rows)
    write_json(out / "meta.json", {**meta, "rows": len(rows), "status": "OK", "primary_file": primary.name})
    write_raw_record(
        source_id=SOURCE_ID,
        ingestion_run_id=run["ingestion_run_id"],
        connector_version=run["connector_version"],
        payload=primary.read_bytes(),
        filename=primary.name,
        source_url=PORTAL,
        dataset_id=DATASET_ID,
    )
    write_manifest(
        out,
        SOURCE_ID,
        [{"file": p.name, "bytes": p.stat().st_size, "study_year": y} for p, y in files],
        extra=meta,
    )
    append_event("document.discovered", {"source": SOURCE_ID, "rows": len(rows)})
    mark_ingested(
        SOURCE_ID,
        run_id=run["ingestion_run_id"],
        counts={"rows": len(rows)},
        dataset_id=DATASET_ID,
        ok=True,
    )
    print(f"OK {SOURCE_ID} rows={len(rows)} -> {out}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
