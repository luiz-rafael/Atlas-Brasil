#!/usr/bin/env python3
"""
Ingest ANA — cobertura/indicadores de água (dados abertos).

Ordem:
  1) ATLAS_ANA_FILE (csv/xlsx)
  2) SNIRH/Atlas Águas — Indicadores_SegurancaHidrica.xlsx (coluna Cobertura)
  3) Portal dadosabertos.ana.gov.br (links CSV) + URLs candidatas
  4) SKIPPED exit 0

Bronze: ana_extract.jsonl
Indicador: ind_ana_cobertura_agua
"""

from __future__ import annotations

import csv
import io
import os
import re
import subprocess
import sys
from pathlib import Path

import httpx

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from pipelines.common import bronze_dir, day_stamp, throttle, utc_now, write_json, write_jsonl  # noqa: E402
from pipelines.registry import mark_ingested, start_run  # noqa: E402

UA = os.getenv("ATLAS_USER_AGENT", "ATLAS-BRASIL-Ingestor/5.1 (+pesquisa ANA)")
PORTAL = "https://dadosabertos.ana.gov.br/"
SNIRH_ATLAS = (
    "https://metadados.snirh.gov.br/geonetwork/srv/api/records/"
    "d77a2d01-0578-4c71-a57e-87f5c565aacf/attachments/"
)

CANDIDATE_URLS = [
    SNIRH_ATLAS + "AtlasAguas_Indicadores_SegurancaHidrica.xlsx",
    SNIRH_ATLAS + "AtlasAguas_MananciaisSistemas.xlsx",
]


def _norm(s: str) -> str:
    import unicodedata

    t = unicodedata.normalize("NFKD", s or "")
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", t.lower())


def download_bytes(url: str, dest: Path) -> dict:
    throttle()
    dest.parent.mkdir(parents=True, exist_ok=True)
    try:
        r = httpx.get(url, headers={"User-Agent": UA}, timeout=180.0, follow_redirects=True)
        if r.status_code == 200 and r.content and len(r.content) > 200:
            dest.write_bytes(r.content)
            return {"url": url, "file": dest.name, "bytes": len(r.content), "via": "httpx"}
        raise RuntimeError(f"HTTP {r.status_code}")
    except Exception as first:
        cmd = [
            "curl.exe" if sys.platform == "win32" else "curl",
            "-L",
            "--retry",
            "2",
            "-A",
            UA,
            "-o",
            str(dest),
            url,
        ]
        try:
            subprocess.run(cmd, check=True, timeout=300)
            if dest.exists() and dest.stat().st_size > 200:
                return {"url": url, "file": dest.name, "bytes": dest.stat().st_size, "via": "curl"}
        except Exception as second:
            return {"url": url, "error": f"{first}; curl: {second}"}
        return {"url": url, "error": str(first)}


def _parse_num(v) -> float | None:
    if v is None:
        return None
    s = str(v).strip().replace("%", "")
    if not s or s.lower() in ("na", "-", "nan"):
        return None
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _pick(fields: list[str], aliases: list[str]) -> str | None:
    norms = {_norm(f): f for f in fields}
    for a in aliases:
        na = _norm(a)
        if na in norms:
            return norms[na]
        for nk, orig in norms.items():
            if na and len(na) >= 5 and na in nk:
                return orig
    return None


def parse_csv_text(text: str, default_year: int | None = None) -> list[dict]:
    delim = ";" if text[:4096].count(";") >= text[:4096].count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    if not reader.fieldnames:
        return []
    fields = list(reader.fieldnames)
    col_ibge = _pick(fields, ["cod_ibge", "codigo_ibge", "codigo ibge", "ibge", "cd_municipio", "co_municipio"])
    col_uf = _pick(fields, ["uf", "sg_uf", "sigla"])
    col_ano = _pick(fields, ["ano", "ano_referencia", "nu_ano"])
    col_cob = _pick(
        fields,
        [
            "cobertura_agua",
            "cobertura",
            "coberturapreenchido",
            "indice_cobertura",
            "atendimento_agua",
            "populacao_atendida_pct",
            "perc_cobertura",
        ],
    )
    if not col_cob:
        return []
    year_default = default_year or int(os.getenv("ATLAS_ANA_YEAR", "2021"))
    now = utc_now()
    out = []
    for row in reader:
        val = _parse_num(row.get(col_cob))
        if val is None:
            continue
        if 0 <= val <= 1:
            val *= 100.0
        year = year_default
        if col_ano and row.get(col_ano):
            ys = re.sub(r"\D", "", str(row[col_ano]))[:4]
            if ys.isdigit():
                year = int(ys)
        code = re.sub(r"\D", "", str(row.get(col_ibge) or "")) if col_ibge else ""
        uf = str(row.get(col_uf) or "").upper()[:2] if col_uf else ""
        if len(code) >= 6:
            code = code[:7] if len(code) >= 7 else code[:6]
            nivel, tid = "MUNICIPALITY", f"mun_{code}"
        elif uf.isalpha():
            nivel, tid, code = "STATE", f"uf_{uf}", uf
        else:
            continue
        out.append(
            {
                "nivel": nivel,
                "territory_id": tid,
                "cod_ibge": code,
                "uf": uf,
                "ano": year,
                "cobertura_agua": val,
                "retrieved_at": now,
                "fonte": "ana",
            }
        )
    return out


def _matrix_to_csv_text(matrix: list[list]) -> str | None:
    """Encontra linha de cabeçalho com IBGE + Cobertura e serializa como CSV."""
    header_idx = None
    for i, row in enumerate(matrix[:30]):
        norms = [_norm(str(c or "")) for c in row]
        if any("ibge" in n or n == "codigoibge" for n in norms) and any(
            "cobertura" in n for n in norms
        ):
            header_idx = i
            break
    if header_idx is None:
        # fallback: qualquer linha com cobertura + uf/municipio
        for i, row in enumerate(matrix[:30]):
            norms = [_norm(str(c or "")) for c in row]
            if any("cobertura" in n for n in norms) and any(
                n in ("uf", "municipio", "codigoibge") or "ibge" in n for n in norms
            ):
                header_idx = i
                break
    if header_idx is None:
        return None
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in matrix[header_idx:]:
        w.writerow(list(r))
    return buf.getvalue()


def parse_file(path: Path, default_year: int | None = None) -> list[dict]:
    if path.suffix.lower() == ".csv":
        return parse_csv_text(path.read_text(encoding="utf-8", errors="replace"), default_year)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            return []
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        all_rows: list[dict] = []
        for name in wb.sheetnames:
            sh = wb[name]
            matrix = [list(r) for r in sh.iter_rows(values_only=True)]
            text = _matrix_to_csv_text(matrix)
            if not text:
                # tenta primeira linha como header
                buf = io.StringIO()
                w = csv.writer(buf)
                for r in matrix:
                    w.writerow(list(r))
                text = buf.getvalue()
            got = parse_csv_text(text, default_year)
            if got:
                all_rows.extend(got)
        # dedupe por território+ano (última ocorrência)
        by_key = {(r["territory_id"], r["ano"]): r for r in all_rows}
        return list(by_key.values())
    return []


def scrape_portal(outdir: Path) -> list[dict]:
    metas = []
    try:
        throttle()
        r = httpx.get(PORTAL, headers={"User-Agent": UA}, timeout=60, follow_redirects=True)
        if r.status_code != 200:
            return [{"page": PORTAL, "error": f"HTTP {r.status_code}"}]
        for href in re.findall(r'href="([^"]+)"', r.text):
            low = href.lower()
            if not any(low.endswith(ext) for ext in (".csv", ".xlsx", ".zip")):
                continue
            if not any(k in low for k in ("agua", "cobertura", "saneamento", "abastec", "atlas")):
                continue
            if href.startswith("/"):
                href = "https://dadosabertos.ana.gov.br" + href
            dest = outdir / Path(href.split("?")[0]).name
            metas.append(download_bytes(href, dest))
    except Exception as e:
        metas.append({"page": PORTAL, "error": str(e)})
    return metas


def main() -> int:
    run = start_run("ana", "ana.cobertura_agua")
    out = bronze_dir("ana")
    raw = out / "raw"
    raw.mkdir(exist_ok=True)
    rows: list[dict] = []
    files_meta: list[dict] = []
    messages: list[str] = []
    # Atlas Águas 2021 (população urbana 2020 / cobertura no diagnóstico)
    atlas_year = int(os.getenv("ATLAS_ANA_YEAR", "2021"))

    local = os.getenv("ATLAS_ANA_FILE")
    if local:
        p = Path(local)
        if p.is_file():
            rows = parse_file(p, atlas_year)
            files_meta.append({"file": str(p), "via": "ATLAS_ANA_FILE", "rows": len(rows)})
        else:
            messages.append(f"ATLAS_ANA_FILE ausente: {local}")

    if not rows:
        for url in CANDIDATE_URLS:
            dest = raw / Path(url.split("?")[0]).name
            meta = download_bytes(url, dest) if not dest.exists() else {
                "url": url,
                "file": dest.name,
                "via": "cache",
                "bytes": dest.stat().st_size,
            }
            files_meta.append(meta)
            if meta.get("error"):
                continue
            # MananciaisSistemas não tem Cobertura — só tenta Indicadores / qualquer com coluna
            got = parse_file(dest, atlas_year)
            if got:
                rows = got
                messages.append(f"SNIRH/Atlas Águas: {dest.name} rows={len(got)}")
                break

    if not rows:
        files_meta.extend(scrape_portal(raw))
        for m in files_meta:
            if m.get("file") and not m.get("error"):
                p = raw / m["file"]
                if p.exists():
                    got = parse_file(p, atlas_year)
                    if got:
                        rows = got
                        break

    if rows:
        write_jsonl(out / "ana_extract.jsonl", rows)
    else:
        messages.append(
            "SKIPPED ana_agua: sem tabular de cobertura. "
            "Defina ATLAS_ANA_FILE=caminho.csv (IBGE/UF + cobertura_agua)."
        )
        print(messages[-1], flush=True)

    write_json(
        out / "ana_meta.json",
        {
            "em": utc_now(),
            "fonte": "ana",
            "rows": len(rows),
            "files": files_meta,
            "messages": messages,
            "ingestion_run_id": run["ingestion_run_id"],
            "day": day_stamp(),
            "portal": PORTAL,
            "snirh": SNIRH_ATLAS,
            "nota": "Cobertura do Atlas Águas 2021 (SNIRH); complementar ao SNIS.",
        },
    )
    mark_ingested(
        "ana",
        run_id=run["ingestion_run_id"],
        counts={"rows": len(rows)},
        ok=len(rows) > 0,
        error=None if rows else (messages[-1] if messages else "sem dados"),
        dataset_id="ana.cobertura_agua",
    )
    print(f"OK ANA: rows={len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
