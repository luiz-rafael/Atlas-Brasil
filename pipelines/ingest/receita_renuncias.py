#!/usr/bin/env python3
"""
Receita Federal — Benefícios / Renúncias fiscais.

Fonte verificada: agregado IRPJ/CSLL/PIS-Imp/COFINS-Imp/IPI-Imp/II (2015–2024).
Nota editorial: renúncia ≠ pagamento público. CNPJ/company_id só com 14 dígitos.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from pipelines.common import (  # noqa: E402
    append_event,
    bronze_dir,
    sha1_bytes,
    utc_now,
    write_json,
    write_jsonl,
    write_manifest,
    write_raw_record,
)
from pipelines.ingest.receita_common import (  # noqa: E402
    company_id_from_cnpj,
    copy_env_file,
    download_to,
    load_dotenv,
    only_cnpj14,
    only_digits,
    scrape_download_links,
    try_urls,
    year_from_header,
)
from pipelines.registry import mark_ingested, start_run  # noqa: E402

SOURCE_ID = "receita_renuncias"
DATASET_ID = "rfb.renuncias"
PORTAL = (
    "https://www.gov.br/receitafederal/pt-br/acesso-a-informacao/"
    "dados-abertos/beneficios-e-renuncias-fiscais"
)
PRIMARY_URL = (
    "https://www.gov.br/receitafederal/dados/"
    "agregado-2015-a-2024-irpj-csll-pis-imp-cofins-imp-ipi-imp-e-ii-5.xlsx/@@download/file"
)
SUBPAGES = [
    PORTAL,
    f"{PORTAL}/renuncias-fiscais-de-tributos-federais",
    f"{PORTAL}/renuncias-fiscais-de-tributos-federais-sobre-importacao",
    f"{PORTAL}/empresas-habilitadas-em-regimes-tributarios-e-aduaneiros-especiais",
]

TRIBUTO_AGREGADO = "IRPJ,CSLL,PIS-Imp,COFINS-Imp,IPI-Imp,II"
BENEFICIO_AGREGADO = "renuncia_tributos_federais_agregado_empresa"
METHODOLOGY = (
    "Agregado RFB a partir de ECF e DW Aduaneiro (extração set/out 2024). "
    "Valores observados de benefício/renúncia por empresa (CNPJ raiz na fonte). "
    "2024 parcial (até 30/06) só tributos de importação. "
    "Renúncia fiscal ≠ pagamento público."
)


def parse_agregado_xlsx(path: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header_idx = None
    year_cols: dict[int, int] = {}
    cnpj_col = None
    razao_col = None
    notes: list[str] = []
    rows_out: list[dict] = []

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        cells = list(row)
        labels = [str(c).strip() if c is not None else "" for c in cells]
        joined = " ".join(labels).strip()
        if i < 5 and joined:
            notes.append(joined[:500])

        if header_idx is None:
            low = [x.lower() for x in labels]
            if any("cnpj" in x for x in low) and any(
                year_from_header(c) for c in cells
            ):
                header_idx = i
                for j, lab in enumerate(labels):
                    ll = lab.lower()
                    if "cnpj" in ll:
                        cnpj_col = j
                    elif "raz" in ll or "social" in ll:
                        razao_col = j
                    else:
                        y = year_from_header(lab)
                        if y and "total" not in ll:
                            year_cols[j] = y
                continue
            continue

        if cnpj_col is None or cnpj_col >= len(cells):
            continue
        raw_cnpj = cells[cnpj_col]
        digits = only_digits(raw_cnpj)
        if not digits:
            continue
        cnpj14 = only_cnpj14(raw_cnpj)
        cnpj_raiz = digits[:8] if len(digits) >= 8 else digits
        razao = None
        if razao_col is not None and razao_col < len(cells) and cells[razao_col]:
            razao = str(cells[razao_col]).strip()

        for col_i, year in year_cols.items():
            if col_i >= len(cells):
                continue
            val = cells[col_i]
            if val is None or val == "":
                continue
            try:
                num = float(val)
            except (TypeError, ValueError):
                continue
            if num == 0:
                continue
            rows_out.append(
                {
                    "year": year,
                    "tributo": TRIBUTO_AGREGADO,
                    "beneficio": BENEFICIO_AGREGADO,
                    "regime": BENEFICIO_AGREGADO,
                    "setor": None,
                    "cnae": None,
                    "valor": num,
                    "value_type": "OBSERVED",
                    "methodology": METHODOLOGY,
                    "methodology_notes": notes[:4],
                    "cnpj": cnpj14,
                    "cnpj_raiz": cnpj_raiz,
                    "company_id": company_id_from_cnpj(cnpj14),
                    "razao_social": razao,
                    "beneficiary_identified": bool(cnpj14),
                    "fonte": SOURCE_ID,
                    "dataset_id": DATASET_ID,
                    "nota": "Renúncia ≠ pagamento. Sem company_id se CNPJ < 14 dígitos.",
                    "source_file": path.name,
                }
            )
    wb.close()
    return rows_out


def discover_extra_urls() -> list[str]:
    found: list[str] = []
    for page in SUBPAGES:
        try:
            found.extend(
                scrape_download_links(
                    page,
                    keywords=("xlsx", "csv", "ods", "renuncia", "beneficio", "agregado"),
                )
            )
        except Exception:
            continue
    # prefer download endpoints
    out: list[str] = []
    for u in found:
        if u.endswith("/view"):
            u = u[: -len("/view")] + "/@@download/file"
        elif "/@@download" not in u and any(
            u.lower().endswith(ext) for ext in (".xlsx", ".csv", ".ods", ".zip")
        ):
            u = u.rstrip("/") + "/@@download/file"
        out.append(u)
    # dedupe
    seen: set[str] = set()
    uniq = []
    for u in out:
        if u not in seen:
            seen.add(u)
            uniq.append(u)
    return uniq


def try_download(out: Path) -> tuple[Path | None, dict]:
    env_path, env_meta = copy_env_file("ATLAS_RENUNCIAS_FILE", out)
    if env_path:
        return env_path, env_meta

    path, meta = download_to(PRIMARY_URL, out / "renuncias_agregado.xlsx", timeout=300.0)
    if path:
        return path, meta

    extras = discover_extra_urls()
    if extras:
        path2, meta2 = try_urls(
            extras[:8],
            out,
            filename="renuncias_extra.xlsx",
            timeout=180.0,
        )
        if path2:
            return path2, {**meta2, "primary_error": meta}
        return None, {"error": meta, "extras": meta2, "candidates": [PRIMARY_URL, *extras[:8]]}
    return None, {"error": meta, "candidates": [PRIMARY_URL]}


def main() -> int:
    load_dotenv()
    try:
        run = start_run(SOURCE_ID, DATASET_ID)
    except Exception as e:
        print(f"fail-soft start_run {SOURCE_ID}: {e}", file=sys.stderr)
        return 0

    out = bronze_dir(SOURCE_ID)
    file_path, meta_dl = try_download(out)
    meta = {
        "fetched_at": utc_now(),
        "portal": PORTAL,
        "ingestion_run_id": run["ingestion_run_id"],
        "nota": "Renúncia fiscal ≠ pagamento público",
        **meta_dl,
    }

    if not file_path:
        meta["status"] = "SKIPPED"
        meta["reason"] = (
            "Download indisponível. Defina ATLAS_RENUNCIAS_FILE com xlsx/csv local."
        )
        write_json(out / "meta.json", meta)
        write_manifest(out, SOURCE_ID, [], extra=meta)
        mark_ingested(SOURCE_ID, run_id=run["ingestion_run_id"], counts={"rows": 0}, dataset_id=DATASET_ID, ok=True)
        print(f"SKIPPED {SOURCE_ID}: sem arquivo. Use ATLAS_RENUNCIAS_FILE.", flush=True)
        return 0

    try:
        rows = parse_agregado_xlsx(file_path)
    except Exception as e:
        meta["status"] = "SKIPPED"
        meta["parse_error"] = str(e)
        write_json(out / "meta.json", meta)
        mark_ingested(SOURCE_ID, run_id=run["ingestion_run_id"], counts={"rows": 0}, dataset_id=DATASET_ID, ok=True)
        print(f"SKIPPED {SOURCE_ID} parse: {e}", flush=True)
        return 0

    write_jsonl(out / "tax_expenditure_extract.jsonl", rows)
    write_json(
        out / "meta.json",
        {
            **meta,
            "rows": len(rows),
            "status": "OK",
            "with_cnpj14": sum(1 for r in rows if r.get("cnpj")),
            "sha1": sha1_bytes(file_path.read_bytes()),
        },
    )
    write_raw_record(
        source_id=SOURCE_ID,
        ingestion_run_id=run["ingestion_run_id"],
        connector_version=run["connector_version"],
        payload=file_path.read_bytes(),
        filename=file_path.name,
        source_url=meta.get("url") or PRIMARY_URL,
        dataset_id=DATASET_ID,
    )
    write_manifest(out, SOURCE_ID, [{"file": file_path.name, "bytes": file_path.stat().st_size}], extra=meta)
    append_event("document.discovered", {"source": SOURCE_ID, "rows": len(rows)})
    mark_ingested(
        SOURCE_ID,
        run_id=run["ingestion_run_id"],
        counts={"rows": len(rows)},
        dataset_id=DATASET_ID,
        ok=True,
    )
    print(f"OK {SOURCE_ID} rows={len(rows)} -> {out}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
