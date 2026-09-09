#!/usr/bin/env python3
"""Peek RFB xlsx structure for parsers."""
from __future__ import annotations

import tempfile
from pathlib import Path

import httpx
import openpyxl

UA = {"User-Agent": "ATLAS-BRASIL-Ingestor/5.1 (+probe)"}


def peek(url: str, name: str, max_sheets: int = 8) -> None:
    print(f"\n==== {name} ====")
    r = httpx.get(url, timeout=180.0, follow_redirects=True, headers=UA)
    print("status", r.status_code, "bytes", len(r.content), "ct", r.headers.get("content-type"))
    if r.status_code != 200 or len(r.content) < 1000:
        return
    p = Path(tempfile.gettempdir()) / name
    p.write_bytes(r.content)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    print("sheets:", wb.sheetnames[:25], "total", len(wb.sheetnames))
    for sn in wb.sheetnames[:max_sheets]:
        ws = wb[sn]
        print(f" -- sheet {sn!r}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 10:
                break
            cells = list(row[:14])
            print("  ", i, cells)
    wb.close()


def main() -> None:
    peek(
        "https://www.gov.br/receitafederal/dados/agregado-2015-a-2024-irpj-csll-pis-imp-cofins-imp-ipi-imp-e-ii-5.xlsx/@@download/file",
        "renuncias.xlsx",
    )
    peek(
        "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/carga-tributaria/tabelas-carga-tributaria-no-brasil-2024/@@download/file",
        "carga2024.xlsx",
        max_sheets=12,
    )


if __name__ == "__main__":
    main()
